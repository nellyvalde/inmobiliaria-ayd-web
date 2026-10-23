"""Genera las tarjetas del inicio y las páginas de detalle desde data/inmuebles.json.

Uso:  python build.py [ruta_del_excel.xlsx]
Cada inmueble lleva un código (101, 102, ...). Para publicar uno nuevo, agrégalo a
data/inmuebles.json (y sus fotos en assets/inmuebles/<codigo>/) y corre este script.
"""
import html
import json
import os
import sys
from urllib.parse import quote

ROOT = os.path.dirname(os.path.abspath(__file__))
WA = '573023487290'
CATEGORIAS = {'proyecto': 'Proyectos', 'lote': 'Lotes', 'casa': 'Casas', 'apartamento': 'Apartamentos', 'arriendo': 'Arriendos'}


def e(t):
    return html.escape(str(t), quote=True)


def wa(texto):
    return f'https://wa.me/{WA}?text={quote(texto)}'


def leer(p):
    with open(os.path.join(ROOT, p), encoding='utf-8') as f:
        return f.read()


def _versionar(s):
    """Agrega ?v=<hash> a imágenes y videos para que el navegador no use una copia vieja guardada."""
    import hashlib
    import re

    def rep(m):
        ruta = m.group(1)
        f = os.path.join(ROOT, ruta.replace('/', os.sep))
        if not os.path.exists(f):
            return m.group(0)
        with open(f, 'rb') as fh:
            h = hashlib.md5(fh.read()).hexdigest()[:8]
        return f'{ruta}?v={h}'
    patron = r"(assets/inmuebles/[^\"'\s?<>]+\.(?:jpg|png|mp4))(?:\?v=[0-9a-f]+)?"
    return re.sub(patron, rep, s)


def escribir(p, s):
    if p.endswith('.html'):
        s = _versionar(s)
    full = os.path.join(ROOT, p)
    os.makedirs(os.path.dirname(full), exist_ok=True)
    with open(full, 'w', encoding='utf-8', newline='\n') as f:
        f.write(s)


def reemplazar(s, ini, fin, nuevo):
    i = s.index(ini) + len(ini)
    j = s.index(fin)
    return s[:i] + '\n' + nuevo + '\n    ' + s[j:]


# ---------- Tarjetas del inicio ----------
def tarjeta(it):
    cod = it['codigo']
    if it.get('foto_portada'):
        media = (f'<div class="h-52 relative"><img src="{e(it["foto_portada"])}" alt="{e(it["titulo"])}" '
                 f'loading="lazy" class="w-full h-full object-cover">')
    else:
        media = f'<div class="{it.get("ph", "ph1")} h-52 relative">'
    badges = (f'<span class="absolute top-3 left-3 bg-vino text-white text-xs font-semibold px-3 py-1 rounded-md">{e(it["negocio"].upper())}</span>'
              f'<span class="absolute top-3 right-3 bg-white/95 text-vino text-xs font-bold px-2.5 py-1 rounded-md">Cód. {cod}</span></div>')
    chips = ''.join(f'<span>{e(c)}</span>' for c in it.get('chips', []))
    nota = f'<p class="text-sm text-tinta/60">{e(it["precio_nota"])}</p>' if it.get('precio_nota') else ''
    if it.get('detalle'):
        href = f'inmueble/{cod}.html'
    else:
        href = wa(it['whatsapp_texto'])
    return f'''    <article data-t="{e(it['categoria'])}" data-cod="{cod}" class="card bg-white rounded-2xl overflow-hidden border border-black/10 flex flex-col">
      {media}{badges}
      <div class="p-6 flex flex-col flex-1">
        <p class="text-xs font-semibold text-rojo uppercase">{e(it['ubicacion_corta'])}</p>
        <h3 class="font-display font-bold text-xl mt-1 text-vino">{e(it['titulo'])}</h3>
        <p class="font-display font-extrabold text-2xl mt-3">{e(it['precio_texto'])}</p>
        {nota}
        <div class="mt-3 flex flex-wrap gap-x-4 gap-y-1 text-sm text-tinta/70">{chips}</div>
        <p class="mt-3 text-sm text-tinta/65">{e(it['resumen'])}</p>
        <a href="{e(href)}" class="mt-auto pt-5"><span class="block text-center rounded-lg border border-rojo text-rojo font-semibold py-2.5 hover:bg-rojo hover:text-white transition">Ver detalles</span></a>
      </div>
    </article>
'''


def hero(items):
    """Carrusel de videos de la portada: una diapositiva por inmueble que tenga 'portada'."""
    slides = [i for i in items if i.get('portada')]
    if not slides:
        return '      <div class="ph1 rounded-3xl h-80 lg:h-[420px] w-full"></div>'
    out = []
    for k, it in enumerate(slides):
        pt = it['portada']
        cod = it['codigo']
        base = f'assets/inmuebles/{cod}/'
        poster = pt.get('poster', 'cover.jpg')
        href = f'inmueble/{cod}.html' if it.get('detalle') else wa(it['whatsapp_texto'])
        out.append(f'''      <article class="hslide absolute inset-0{' is-active' if k == 0 else ''}" data-src="{base}{e(pt['video'])}" data-start="{pt.get('inicio', 0)}" data-end="{pt.get('fin', 15)}">
        <video muted playsinline preload="none" poster="{base}{e(poster)}" class="w-full h-full object-cover"></video>
        <div class="absolute inset-0 bg-gradient-to-b from-black/75 via-black/15 to-transparent"></div>
        <div class="absolute left-5 right-5 top-5 text-white">
          <p class="text-xs font-semibold uppercase tracking-wide text-white/80">Cód. {cod} · {e(pt.get('lugar', it['ubicacion_corta']))}</p>
          <h3 class="mt-1 font-display font-extrabold text-xl md:text-2xl leading-tight">{e(pt.get('titulo', it['titulo']))}</h3>
          <p class="mt-1 text-sm text-white/90">{e(pt.get('linea', ''))}</p>
          <a href="{e(href)}" class="inline-block mt-3 rounded-lg bg-white text-vino text-sm font-semibold px-4 py-2 hover:bg-cielo transition">Ver {'proyecto' if it['categoria'] == 'proyecto' else 'inmueble'}</a>
        </div>
      </article>''')
    return f'''      <div id="hero" class="relative rounded-3xl overflow-hidden h-80 lg:h-[420px] w-full bg-vino" role="region" aria-roledescription="carrusel" aria-label="Proyectos destacados">
{chr(10).join(out)}
        <div id="hdots" class="absolute bottom-5 left-5 flex items-center gap-2"></div>
        <button type="button" id="hpause" class="absolute bottom-3 right-4 h-9 w-9 rounded-full bg-black/40 text-white text-xs hover:bg-black/60" aria-label="Pausar">❚❚</button>
      </div>'''


def inicio(items):
    s = leer('index.html')
    s = reemplazar(s, '<!-- HERO:START -->', '<!-- HERO:END -->', hero(items))
    cats = []
    for it in items:
        if it['categoria'] not in cats:
            cats.append(it['categoria'])
    tabs = '      <button data-f="todos" class="tab-on rounded-lg px-4 py-2 bg-cielo text-vino">Todos</button>\n'
    for c in cats:
        tabs += f'      <button data-f="{c}" class="rounded-lg px-4 py-2 bg-cielo text-vino">{CATEGORIAS.get(c, c.title())}</button>\n'
    s = reemplazar(s, '<!-- TABS:START -->', '<!-- TABS:END -->', tabs.rstrip('\n'))
    s = reemplazar(s, '<!-- CARDS:START -->', '<!-- CARDS:END -->', '\n'.join(tarjeta(i) for i in items).rstrip('\n'))
    escribir('index.html', s)


# ---------- Página de detalle ----------
def partes_del_sitio():
    s = leer('index.html')
    head = s[:s.index('</head>')]
    chrome_ini = s.index('<!-- Franja superior -->')
    chrome_fin = s.index('<!-- Portada -->')
    top = s[chrome_ini:chrome_fin]
    foot_ini = s.index('<footer')
    foot_fin = s.index('<script', foot_ini)
    foot = s[foot_ini:foot_fin]
    # rutas relativas desde /inmueble/
    top = top.replace('href="#', 'href="../index.html#').replace('src="logo.svg"', 'src="../logo.svg"')
    head = head.replace('href="logo.svg"', 'href="../logo.svg"')
    return head, top, foot


def lista(titulo, puntos):
    if not puntos:
        return ''
    li = ''.join(f'<li class="flex gap-3"><span class="text-rojo font-bold">&#10003;</span><span>{e(p)}</span></li>' for p in puntos)
    return f'''<div class="rounded-2xl border border-black/10 p-6">
        <h3 class="font-display font-bold text-lg text-vino">{e(titulo)}</h3>
        <ul class="mt-4 space-y-2 text-tinta/80">{li}</ul>
      </div>'''


def dato(titulo, valor, extra=''):
    return f'''<div class="rounded-2xl bg-suave p-5">
        <p class="text-xs font-semibold text-rojo uppercase tracking-wide">{e(titulo)}</p>
        <p class="mt-1 font-display font-extrabold text-xl text-tinta">{e(valor)}</p>
        {f'<p class="text-sm text-tinta/60">{e(extra)}</p>' if extra else ''}
      </div>'''


def detalle(it, partes):
    head, top, foot = partes
    cod = it['codigo']
    base = f'../assets/inmuebles/{cod}/'
    fotos = it.get('fotos', [])
    mapas = it.get('mapas', [])
    thumbs = ''.join(
        f'<button type="button" class="thumb rounded-lg overflow-hidden border-2 border-transparent aspect-[4/3] bg-black/5" data-i="{k}">'
        f'<img src="{base}{e(f["f"])}" alt="{e(f["alt"])}" loading="lazy" class="w-full h-full object-cover"></button>'
        for k, f in enumerate(fotos))
    gal_json = json.dumps([{'src': base + f['f'], 'alt': f['alt']} for f in fotos], ensure_ascii=False)

    datos = ''
    datos += dato('Área', it['area_texto'])
    datos += dato('Precio', it['precio_texto'], it.get('precio_nota', ''))
    if it.get('valor_total_aprox'):
        datos += dato('Valor aproximado', it['valor_total_aprox'], f'{it["area_texto"]} × {it["precio_texto"].replace(" por m²", "")}')
    for d in it.get('datos_extra', []):
        datos += dato(d['t'], d['v'], d.get('x', ''))
    if it.get('permuta'):
        datos += dato('Permuta', 'Sí', it['permuta'])

    cols = 'lg:grid-cols-5' if datos.count('rounded-2xl bg-suave') >= 5 else 'lg:grid-cols-4'
    desc = ''.join(f'<p class="mt-3 text-tinta/75">{e(p)}</p>' for p in it.get('descripcion', []))

    mapas_html = ''.join(
        f'''<figure><button type="button" class="zoom block w-full rounded-2xl overflow-hidden border border-black/10 bg-black/5" data-src="{base}{e(m['f'])}" data-alt="{e(m['texto'])}"><img src="{base}{e(m['f'])}" alt="{e(m['texto'])}" loading="lazy" class="w-full {m.get('aspect', 'aspect-[3/4] object-cover')}"></button>
          <figcaption class="mt-2 text-sm text-tinta/65">{e(m['texto'])}</figcaption></figure>'''
        for m in mapas)

    poster_img = it.get('poster') or (fotos[0]['f'] if fotos else 'cover.jpg')
    videos = it.get('videos') or ([it['video']] if it.get('video') else [])
    video = ''
    if videos:
        vids = ''.join(
            f'''<video class="w-full max-h-[70vh] rounded-2xl bg-black" controls preload="metadata" playsinline poster="{base}{poster_img}">
        <source src="{base}{e(v)}" type="video/mp4">
      </video>''' for v in videos)
        video = f'''<section class="mt-10">
      <h2 class="font-display font-extrabold text-2xl text-vino">{'Videos' if len(videos) > 1 else 'Video'}</h2>
      <div class="mt-4 grid gap-5 {'md:grid-cols-3' if len(videos) > 2 else ('md:grid-cols-2' if len(videos) > 1 else '')}">{vids}</div>
    </section>'''

    mapas_sec = ''
    if mapas:
        mapas_sec = f'''<section class="mt-14">
    <h2 class="font-display font-extrabold text-2xl text-vino">Mapa y ubicación</h2>
    <div class="mt-4 grid {'md:grid-cols-3' if len(mapas) > 1 else ''} gap-5">{mapas_html}</div>
  </section>'''

    listas = ''.join([
        lista(it.get('servicios_titulo', 'Servicios'), it.get('servicios')),
        lista('Tipologías y precios', it.get('tipologias')),
        lista(it.get('incluye_titulo', 'Incluye'), it.get('incluye')),
        lista('Razones para elegirnos', it.get('razones')),
        lista('Condiciones', it.get('condiciones')),
        lista('Ubicación', it.get('ubicacion_puntos')),
    ])
    una_foto = len(fotos) == 1
    grid_gal = '' if una_foto else 'lg:grid-cols-[1fr_300px]'
    main_aspect = 'p-3' if una_foto else 'aspect-[4/3]'
    img_cls = 'max-h-[78vh] w-auto mx-auto rounded-xl' if una_foto else 'w-full h-full object-contain'
    thumbs_html = '' if una_foto else f'<div class="grid grid-cols-4 lg:grid-cols-3 gap-2 content-start max-h-[520px] overflow-y-auto pr-1" id="thumbs">{thumbs}</div>'

    wa_link = wa(it['whatsapp_texto'])
    titulo_pag = f'Cód. {cod} · {it["titulo"]} | Inmobiliaria A&amp;D'

    head = head.replace(head[head.index('<title>'):head.index('</title>') + 8], f'<title>{titulo_pag}</title>')
    import re
    og_titulo = f"Cód. {cod} · {it['titulo']} | Inmobiliaria A&amp;D"
    og_desc = e(it['precio_texto'] + ' · ' + it['resumen'])
    portada = it.get('foto_portada', '')
    head = re.sub(r'<meta name="description" content="[^"]*">', f'<meta name="description" content="{og_desc}">', head)
    head = re.sub(r'<meta property="og:url" content="[^"]*">', f'<meta property="og:url" content="https://inmobiliariaayd.online/inmueble/{cod}.html">', head)
    head = re.sub(r'<meta property="og:title" content="[^"]*">', f'<meta property="og:title" content="{og_titulo}">', head)
    head = re.sub(r'<meta property="og:description" content="[^"]*">', f'<meta property="og:description" content="{og_desc}">', head)
    if portada:
        head = re.sub(r'<meta property="og:image" content="[^"]*">', f'<meta property="og:image" content="https://inmobiliariaayd.online/{portada}">', head)
        head = re.sub(r'\s*<meta property="og:image:(width|height)" content="[^"]*">', '', head)
    return f'''{head}
</head>
<body class="antialiased">

{top}
<main class="max-w-6xl mx-auto px-5 py-10">
  <nav class="text-sm text-tinta/60"><a href="../index.html" class="hover:text-rojo">Inicio</a> / <a href="../index.html#inmuebles" class="hover:text-rojo">Inmuebles</a> / <span>Cód. {cod}</span></nav>

  <div class="mt-5 flex flex-wrap items-start justify-between gap-4">
    <div>
      <div class="flex flex-wrap gap-2 text-xs font-semibold">
        <span class="bg-vino text-white px-3 py-1 rounded-md">{e(it['negocio'].upper())}</span>
        <span class="bg-cielo text-vino px-3 py-1 rounded-md">Cód. {cod}</span>
        <span class="bg-cielo text-vino px-3 py-1 rounded-md">{e(it.get('estado', ''))}</span>
      </div>
      <h1 class="mt-3 font-display font-extrabold text-3xl md:text-4xl text-vino">{e(it['titulo'])}</h1>
      <p class="mt-2 text-tinta/70">{e(it['referencia'])}</p>
    </div>
    <a href="{wa_link}" class="rounded-lg bg-rojo text-white font-semibold px-6 py-3.5 hover:bg-vino transition">Consultar por WhatsApp</a>
  </div>

  <section class="mt-8 grid {grid_gal} gap-4">
    <button type="button" id="main" class="relative rounded-2xl overflow-hidden bg-black/5 {main_aspect} w-full">
      <img id="mainimg" src="{base}{fotos[0]['f']}" alt="{e(fotos[0]['alt'])}" class="{img_cls}">
      <span id="counter" class="absolute bottom-3 right-3 bg-black/60 text-white text-xs px-3 py-1 rounded-full"></span>
    </button>
    {thumbs_html}
  </section>

  <section class="mt-10 grid sm:grid-cols-2 {cols} gap-4">
    {datos}
  </section>

  {video}

  <section class="mt-10">
    <h2 class="font-display font-extrabold text-2xl text-vino">Descripción</h2>
    {desc}
  </section>

  <section class="mt-10 grid md:grid-cols-2 lg:grid-cols-3 gap-6">
    {listas}
  </section>

  {mapas_sec}

  <section class="mt-14 rounded-3xl bg-rojo text-white px-8 py-12 text-center">
    <h2 class="font-display font-extrabold text-3xl max-w-2xl mx-auto">¿Te interesa el inmueble {cod}?</h2>
    <p class="mt-3 text-white/85">{e(it.get('permuta', 'Escríbenos y coordinamos una visita.'))}</p>
    <a href="{wa_link}" class="inline-block mt-7 rounded-lg bg-white text-vino font-semibold px-8 py-4 hover:bg-cielo transition">Escribir por WhatsApp</a>
  </section>
</main>

{foot}
<div id="lb" class="hidden fixed inset-0 z-[60] bg-black/90 items-center justify-center p-4">
  <button type="button" id="lbclose" class="absolute top-4 right-5 text-white text-4xl leading-none" aria-label="Cerrar">&times;</button>
  <button type="button" id="lbprev" class="absolute left-3 top-1/2 text-white text-5xl leading-none px-3" aria-label="Anterior">&#8249;</button>
  <img id="lbimg" src="" alt="" class="max-h-[90vh] max-w-full object-contain">
  <button type="button" id="lbnext" class="absolute right-3 top-1/2 text-white text-5xl leading-none px-3" aria-label="Siguiente">&#8250;</button>
</div>

<script>
const G = {gal_json};
let cur = 0, lista = G, lbOpen = false;
const mainimg = document.getElementById('mainimg'), counter = document.getElementById('counter');
const thumbs = document.querySelectorAll('.thumb');
function show(i) {{
  cur = (i + G.length) % G.length;
  mainimg.src = G[cur].src; mainimg.alt = G[cur].alt;
  counter.textContent = (cur + 1) + ' / ' + G.length;
  thumbs.forEach((t, k) => t.classList.toggle('border-rojo', k === cur));
}}
thumbs.forEach(t => t.addEventListener('click', () => show(+t.dataset.i)));
const lb = document.getElementById('lb'), lbimg = document.getElementById('lbimg');
function openLb(list, i) {{
  lista = list; cur = i; lbOpen = true;
  lbimg.src = lista[cur].src; lbimg.alt = lista[cur].alt;
  lb.classList.remove('hidden'); lb.classList.add('flex');
  const multi = lista.length > 1;
  document.getElementById('lbprev').style.display = multi ? '' : 'none';
  document.getElementById('lbnext').style.display = multi ? '' : 'none';
}}
function stepLb(d) {{ cur = (cur + d + lista.length) % lista.length; lbimg.src = lista[cur].src; lbimg.alt = lista[cur].alt; }}
function closeLb() {{ lbOpen = false; lb.classList.add('hidden'); lb.classList.remove('flex'); }}
document.getElementById('main').addEventListener('click', () => openLb(G, cur));
document.querySelectorAll('.zoom').forEach(z => z.addEventListener('click', () => openLb([{{src: z.dataset.src, alt: z.dataset.alt}}], 0)));
document.getElementById('lbclose').addEventListener('click', closeLb);
document.getElementById('lbprev').addEventListener('click', () => stepLb(-1));
document.getElementById('lbnext').addEventListener('click', () => stepLb(1));
lb.addEventListener('click', ev => {{ if (ev.target === lb) closeLb(); }});
document.addEventListener('keydown', ev => {{
  if (!lbOpen) return;
  if (ev.key === 'Escape') closeLb();
  if (ev.key === 'ArrowLeft') stepLb(-1);
  if (ev.key === 'ArrowRight') stepLb(1);
}});
show(0);
</script>
</body>
</html>
'''


# ---------- Excel ----------
def _dividir(texto):
    """Divide el texto de pendientes en preguntas sueltas, sin cortar lo que va entre paréntesis."""
    partes, actual, prof = [], '', 0
    for k, ch in enumerate(texto):
        if ch in '([':
            prof += 1
        elif ch in ')]':
            prof = max(0, prof - 1)
        if prof == 0 and (ch == ';' or (ch == '.' and texto[k + 1:k + 2] == ' ')):
            partes.append(actual)
            actual = ''
        else:
            actual += ch
    partes.append(actual)
    return [p.strip() for p in partes if p.strip()]


def excel(items, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Inventario'
    cols = ['Código', 'Negocio', 'Tipo', 'Título', 'Ubicación', 'Referencia de ubicación', 'Área', 'Precio',
            'Nota de precio', 'Valor total aprox.', 'Permuta', 'Separación', 'Financiación', 'Servicios',
            'Incluye', 'Documentos / legal', 'Descripción', 'Estado', 'Fecha', 'Fotos', 'Video',
            'Enlace en la página', 'Por confirmar con Deisy']
    ws.append(cols)
    for it in items:
        ws.append([
            it['codigo'], it['negocio'], it.get('tipo', ''), it['titulo'], it['ubicacion_corta'], it['referencia'],
            it['area_texto'], it['precio_texto'], it.get('precio_nota', ''), it.get('valor_total_aprox', ''),
            it.get('permuta', ''), it.get('separacion', ''), it.get('financiacion', ''),
            '; '.join(it.get('servicios', [])), '; '.join(it.get('incluye', [])), it.get('documentos', ''),
            ' '.join(it.get('descripcion', [])) or it.get('resumen', ''), it.get('estado', ''), it.get('fecha', ''),
            len(it.get('fotos', [])), 'Sí' if it.get('video') else '',
            f'https://inmobiliariaayd.online/inmueble/{it["codigo"]}.html' if it.get('detalle') else 'Solo tarjeta (WhatsApp)',
            it.get('por_confirmar', '')])
    head_fill = PatternFill('solid', fgColor='8B0F1F')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical='center')
    anchos = [9, 9, 22, 30, 30, 40, 14, 18, 22, 18, 30, 14, 30, 45, 28, 30, 60, 13, 12, 7, 7, 38, 50]
    for k, w in enumerate(anchos):
        ws.column_dimensions[chr(65 + k)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'E2'

    # Hoja de preguntas pendientes para Deisy (una fila por pregunta)
    wq = wb.create_sheet('Preguntas para Deisy')
    wq.append(['Código', 'Inmueble', 'Pregunta o dato por confirmar', 'Respuesta de Deisy', 'Resuelto'])
    generales = [
        'Nombre definitivo de la inmobiliaria (cuando quede registrada en Cámara de Comercio) y si tiene logo propio',
        'Foto de Deisy o del equipo para la sección "Quiénes somos"',
        'Autorización de las constructoras para ofrecer sus proyectos (Torres de Pamplonita, Arrayanes de Peñalisa, Santo Domingo, Acqua Viva, La Colina, Rincón del Palmar, Santorini, Sky 360)',
        'Autorización de LQN y VehiCredi para nombrarlos en la página, y si las condiciones del servicio de acompañamiento en financiación están correctas',
        'Fotos de casas prefabricadas ya construidas (y de los otros servicios) para ponerlas en la página; y si los aliados que construyen están de acuerdo con que se mencione el servicio',
        'Si los servicios de la página están bien: arriendos, venta, construcción y remodelaciones, fincas y lotes, trámite de licencias y persianas y cortinas',
    ]
    for q in generales:
        wq.append(['General', 'Toda la página', q, '', ''])
    for it in items:
        pc = it.get('por_confirmar', '')
        partes = _dividir(pc)
        for q in partes:
            wq.append([it['codigo'], it['titulo'], q.rstrip('.'), '', ''])
    for c in wq[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for k, w in enumerate([10, 40, 90, 45, 10]):
        wq.column_dimensions[chr(65 + k)].width = w
    amarillo = PatternFill('solid', fgColor='FFF2CC')
    for row in wq.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
        row[3].fill = amarillo
    wq.freeze_panes = 'A2'
    wq.auto_filter.ref = f'A1:E{wq.max_row}'
    wb.save(ruta)


def ficha_dinamica():
    """Página ficha.html: muestra los inmuebles creados desde el panel (se leen de Supabase)."""
    s = leer('index.html')
    head = s[:s.index('</head>')]
    i = s.index('<!-- Franja superior -->'); j = s.index('<!-- Portada -->')
    top = s[i:j]
    k = s.index('<footer'); l = s.index('<script', k)
    foot = s[k:l]
    head = head.replace(head[head.index('<title>'):head.index('</title>') + 8], '<title>Inmueble | Inmobiliaria A&amp;D</title>')
    import re
    head = re.sub(r'\s*<meta property="og:[^>]*>', '', head)
    head = re.sub(r'\s*<meta name="twitter:[^>]*>', '', head)
    head = re.sub(r'<meta name="description" content="[^"]*">', '<meta name="description" content="Ficha de inmueble de Inmobiliaria A&amp;D">', head)
    html_ = f'''{head}
</head>
<body class="antialiased">

{top}
<main id="ficha" class="max-w-6xl mx-auto px-5 py-10"><p class="py-24 text-center text-tinta/60">Cargando inmueble…</p></main>

{foot}
<div id="lb" class="hidden fixed inset-0 z-[60] bg-black/90 items-center justify-center p-4">
  <button type="button" id="lbclose" class="absolute top-4 right-5 text-white text-4xl leading-none" aria-label="Cerrar">&times;</button>
  <button type="button" id="lbprev" class="absolute left-3 top-1/2 text-white text-5xl leading-none px-3" aria-label="Anterior">&#8249;</button>
  <img id="lbimg" src="" alt="" class="max-h-[90vh] max-w-full object-contain">
  <button type="button" id="lbnext" class="absolute right-3 top-1/2 text-white text-5xl leading-none px-3" aria-label="Siguiente">&#8250;</button>
</div>
<script src="config.js"></script>
<script src="ficha.js"></script>
</body>
</html>
'''
    escribir('ficha.html', html_)


def publicos(items):
    """Lista mínima de los inmuebles estáticos, para que el panel los traiga y permita ocultarlos o marcarlos como vendidos."""
    out = [{'codigo': i['codigo'], 'titulo': i['titulo'], 'categoria': i['categoria'], 'negocio': i['negocio'], 'estado': i.get('estado', 'Disponible'),
            'precio_texto': i['precio_texto'], 'precio_nota': i.get('precio_nota', ''), 'ubicacion_corta': i['ubicacion_corta']} for i in items]
    escribir('inmuebles-publicos.json', json.dumps(out, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    with open(os.path.join(ROOT, 'data', 'inmuebles.json'), encoding='utf-8') as f:
        items = json.load(f)
    codigos = [i['codigo'] for i in items]
    assert len(codigos) == len(set(codigos)), 'Código repetido'
    inicio(items)
    ficha_dinamica()
    publicos(items)
    partes = partes_del_sitio()
    for it in items:
        if it.get('detalle'):
            escribir(f'inmueble/{it["codigo"]}.html', detalle(it, partes))
    if len(sys.argv) > 1:
        excel(items, sys.argv[1])
    print('Listo:', len(items), 'inmuebles')
